"""
test_picos_schema.py — test the PICOS → Schema Generator.

Tests:
  - DeterministicSchemaGenerator with injected fields
  - generate_project_config produces a loadable LitExtract config
  - Merged config has correct field counts (base + PICOS)
  - Generated config passes LitExtract's load_project validation
  - PICOS fields appear in schema.yaml, codebook.xlsx, prompts, and validation rules
"""
import os
import sys
import json
import yaml
import pytest

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO not in sys.path:
    sys.path.insert(0, REPO)

from radextract_core.radextract.picos_schema_generator import (
    PICOS, GeneratedField, DeterministicSchemaGenerator,
    generate_project_config,
)
import radextract_core.litextract as lx


# --------------------------------------------------------------------------- #
# Fixtures
# --------------------------------------------------------------------------- #
BASE_CONFIG = os.path.join(REPO, "configs", "radiology_base")
OUTPUT_DIR = os.path.join(REPO, "configs", "_test_generated")

PICOS_GLIOMA = PICOS(
    population="patients with suspected glioma undergoing MRI",
    intervention="DSC perfusion MRI",
    comparator="histopathology (WHO 2021 grading)",
    outcome="diagnostic accuracy for IDH mutation status",
    study_design="diagnostic accuracy studies",
)

PICOS_FIELDS = [
    GeneratedField(
        sheet="Cohort_Level", order=17, name="IDH Mutation Status",
        field_type="enum", allowed_values=["mutant", "wildtype", "unknown", "NR"],
        description="IDH mutation status of the cohort",
        why="Essential for assessing diagnostic accuracy of IDH prediction",
    ),
    GeneratedField(
        sheet="Cohort_Level", order=18, name="WHO Grade",
        field_type="enum", allowed_values=["2", "3", "4", "unknown", "NR"],
        description="WHO grade of gliomas in the cohort",
        why="Tumor grade affects perfusion metrics",
    ),
    GeneratedField(
        sheet="Protocol_Level", order=25, name="DSC Sequence Parameters",
        field_type="free_text", allowed_values=[],
        description="DSC perfusion sequence parameters (GRE vs SE, flip angle, etc.)",
        why="DSC acquisition type affects CBV measurement",
    ),
    GeneratedField(
        sheet="Protocol_Level", order=26, name="CBV Measurement Method",
        field_type="enum",
        allowed_values=["semi_quantitative", "quantitative", "other", "NR"],
        description="Method used to measure CBV",
        why="Measurement method affects comparability across studies",
    ),
    GeneratedField(
        sheet="Outcome_Level", order=27, name="IDH Detection Sensitivity (%)",
        field_type="percent", allowed_values=[],
        range_min=0, range_max=100,
        description="Sensitivity for detecting IDH mutation status",
        why="Primary outcome metric for this PICOS",
    ),
    GeneratedField(
        sheet="Outcome_Level", order=28, name="IDH Detection Specificity (%)",
        field_type="percent", allowed_values=[],
        range_min=0, range_max=100,
        description="Specificity for detecting IDH mutation status",
        why="Primary outcome metric for this PICOS",
    ),
    GeneratedField(
        sheet="Outcome_Level", order=29, name="IDH Detection AUC",
        field_type="float", allowed_values=[],
        range_min=0, range_max=1,
        description="AUC for IDH mutation detection",
        why="Overall diagnostic performance metric",
    ),
    GeneratedField(
        sheet="Outcome_Level", order=30, name="Optimal CBV Cutoff",
        field_type="float", allowed_values=[],
        description="Optimal CBV threshold for IDH prediction",
        why="Threshold value for clinical translation",
    ),
]


# --------------------------------------------------------------------------- #
# Tests
# --------------------------------------------------------------------------- #
class TestDeterministicSchemaGenerator:
    def test_generate_with_static_fields(self):
        gen = DeterministicSchemaGenerator(fields=PICOS_FIELDS)
        result = gen.generate_fields(PICOS_GLIOMA)
        assert len(result) == 8
        assert result[0].name == "IDH Mutation Status"
        assert result[0].sheet == "Cohort_Level"

    def test_generate_with_fn(self):
        def fn(picos):
            return [
                GeneratedField(sheet="Outcome_Level", order=27,
                               name="Test Field", field_type="free_text",
                               allowed_values=[], description="test", why="test"),
            ]
        gen = DeterministicSchemaGenerator(fields_fn=fn)
        result = gen.generate_fields(PICOS_GLIOMA)
        assert len(result) == 1
        assert result[0].name == "Test Field"


class TestGenerateProjectConfig:
    @pytest.fixture(scope="class")
    def generated_config(self):
        """Generate a project config once for all tests in this class."""
        gen = DeterministicSchemaGenerator(fields=PICOS_FIELDS)
        result = generate_project_config(
            PICOS_GLIOMA, gen, BASE_CONFIG, OUTPUT_DIR,
            project_slug="test_glioma_idh",
        )
        return result

    def test_config_dir_created(self, generated_config):
        assert os.path.isdir(generated_config["config_dir"])

    def test_field_counts(self, generated_config):
        assert generated_config["n_base_fields"] == 107
        assert generated_config["n_picos_fields"] == 8
        assert generated_config["n_total_fields"] == 115

    def test_project_yaml_exists(self, generated_config):
        path = os.path.join(generated_config["config_dir"], "project.yaml")
        assert os.path.exists(path)
        with open(path) as f:
            meta = yaml.safe_load(f)
        assert meta["name"] == "test_glioma_idh"

    def test_picos_yaml_exists(self, generated_config):
        path = os.path.join(generated_config["config_dir"], "picos.yaml")
        assert os.path.exists(path)
        with open(path) as f:
            picos = yaml.safe_load(f)
        assert picos["population"] == PICOS_GLIOMA.population
        assert picos["intervention"] == PICOS_GLIOMA.intervention

    def test_schema_yaml_has_picos_fields(self, generated_config):
        path = os.path.join(generated_config["config_dir"], "schema.yaml")
        with open(path) as f:
            schema = yaml.safe_load(f)
        ft = schema.get("field_types", {})
        # Check a PICOS field is present
        assert "Cohort_Level::IDH Mutation Status" in ft
        assert ft["Cohort_Level::IDH Mutation Status"]["type"] == "enum"
        assert "Outcome_Level::IDH Detection AUC" in ft
        assert ft["Outcome_Level::IDH Detection AUC"]["type"] == "float"

    def test_validation_rules_has_picos_rules(self, generated_config):
        path = os.path.join(generated_config["config_dir"], "validation_rules.yaml")
        with open(path) as f:
            rules_doc = yaml.safe_load(f)
        rules = rules_doc.get("rules", [])
        # Check that PICOS-generated range rules are present
        picos_rules = [r for r in rules if r.get("id", "").startswith("picos_")]
        assert len(picos_rules) >= 3  # at least the 3 range rules

    def test_prompts_has_picos_prompts(self, generated_config):
        path = os.path.join(generated_config["config_dir"], "prompts", "field_prompts.json")
        with open(path) as f:
            prompts = json.load(f)
        names = [p["name"] for p in prompts]
        assert "IDH Mutation Status" in names
        assert "IDH Detection AUC" in names
        assert len(prompts) == 115  # 107 base + 8 PICOS

    def test_codebook_has_picos_fields(self, generated_config):
        path = os.path.join(generated_config["config_dir"], "codebook.xlsx")
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        cb = wb["Codebook"]
        names = set()
        for r in range(2, cb.max_row + 1):
            name = cb.cell(r, 3).value
            if name:
                names.add(str(name).strip())
        wb.close()
        assert "IDH Mutation Status" in names
        assert "IDH Detection AUC" in names

    def test_config_loads_with_litextract(self, generated_config):
        """The generated config must load with LitExtract's load_project."""
        project = lx.load_project(generated_config["config_dir"])
        assert project.name == "test_glioma_idh"
        assert len(project.data_sheets()) == 5
        # Check that PICOS fields are in the sheet_order
        cohort_fields = [name for _, name in project.sheet_order.get("Cohort_Level", [])]
        assert "IDH Mutation Status" in cohort_fields

    def test_config_validates_records(self, generated_config):
        """The generated config should validate records without errors."""
        project = lx.load_project(generated_config["config_dir"])
        # Create a minimal valid record
        records = {
            "Study_Level": [{
                "Study ID": "Test2024", "First Author": "Test",
                "Publication Year": "2024", "Journal": "Test Journal",
                "Country": "USA", "Study Design": "diagnostic_accuracy",
                "Study Setting": "single_center",
                "Conflict of Interest": "none_declared",
                "Registration": "none",
            }],
            "Cohort_Level": [{
                "Study ID": "Test2024", "Cohort ID": "C1",
                "Cohort Role": "disease", "Disease/Condition": "glioma",
                "Sample Size (Enrolled)": 50, "Age Group": "adult",
                "Clinical Indication": "glioma evaluation",
                "IDH Mutation Status": "mutant",
            }],
            "Protocol_Level": [{
                "Cohort ID": "C1", "Protocol ID": "P1",
                "Modality": "MRI", "Field Strength (T)": "3.0",
                "Contrast Agent": "gadolinium", "Contrast Dose": "0.1 mmol/kg",
            }],
            "Outcome_Level": [{
                "Protocol ID": "P1", "Outcome ID": "O1",
                "Result Type": "diagnostic_accuracy",
                "Specific Outcome Label": "IDH detection accuracy",
                "Reference Standard": "histopathology",
                "ROI/Region Category": "lesion", "N Analyzed": 50,
                "Sensitivity (%)": 85.0, "Specificity (%)": 90.0,
                "Unit (as reported)": "percent",
                "Standardized Unit Family": "percent",
                "IDH Detection Sensitivity (%)": 85.0,
                "IDH Detection Specificity (%)": 90.0,
                "IDH Detection AUC": 0.92,
            }],
            "Quality_Assessment": [{
                "Study ID": "Test2024", "QA ID": "QA1",
                "Assessment Tool": "QUADAS_2",
                "D1: Patient Selection — Risk of Bias": "low",
                "D2: Index Test — Risk of Bias": "low",
                "D3: Reference Standard — Risk of Bias": "low",
                "D4: Flow and Timing — Risk of Bias": "low",
                "Overall Risk of Bias": "low",
                "Overall Applicability": "low",
            }],
        }
        issues, summary = lx.validate(records, project)
        errors = [i for i in issues if i["severity"] == "error"]
        assert len(errors) == 0, f"Unexpected validation errors: {errors}"
        assert summary["ok"] is True

    def test_cleanup(self, generated_config):
        """Clean up generated config after tests."""
        import shutil
        if os.path.isdir(generated_config["config_dir"]):
            shutil.rmtree(generated_config["config_dir"])
