"""
codebook.py — generic codebook loader.

The codebook workbook (a sheet, by default named ``Codebook``) is the single
source of truth for field order, formats, controlled vocabularies, and required
flags. This is a generalization of the proven CVR ``load_codebook`` logic: the
*data-sheet names* are no longer hardcoded — they come from the codebook itself
(and are cross-checked against ``schema.yaml``).

Codebook column layout (1-indexed), matching the existing CVR workbook:
    1 Worksheet | 2 Column Order | 3 Final Column Name | 4 Why | 5 What to enter
    6 Expected Format | 7 Allowed/Standardized Entries | 8 Required? | 9 Notes
Column indices are configurable via ``col_map`` for other projects.
"""
from __future__ import annotations
import re
from typing import Optional

FREE_TEXT_SENTINEL = "Unique free text"
DEFAULT_COL_MAP = {
    "sheet": 1, "order": 2, "name": 3, "why": 4, "enter": 5,
    "format": 6, "allowed": 7, "required": 8, "notes": 9,
}
_FREE_MARKERS = ("free text", "integer", "number", "text", "use the review",
                 "4-digit", "0-100", "existing", "percentage")
# Numeric range pattern: "0-1", "1-100000", "0.0-1.0", "-5-50", etc.
_NUMERIC_RANGE_RE = re.compile(r'^-?\d+(?:\.\d+)?\s*[-–]\s*-?\d+(?:\.\d+)?$')


def split_allowed(s) -> list:
    """Parse an 'Allowed/Standardized Entries' cell into a token list.

    Controlled vocabularies are ';'-separated. Free-text / numeric / instruction
    cells return [] (no closed vocabulary).
    """
    if s is None:
        return []
    s = str(s).strip()
    low = s.lower()
    if any(low.startswith(m) for m in _FREE_MARKERS):
        return []
    # Numeric ranges (e.g. "0-1", "1-100000", "0.0-1.0") are not controlled vocab
    if _NUMERIC_RANGE_RE.match(low):
        return []
    return [p.strip() for p in s.split(";") if p.strip()]


def load_codebook(workbook_path: str, *, codebook_sheet: str = "Codebook",
                  col_map: Optional[dict] = None):
    """Return (codebook, sheet_order, controlled_vocab).

    codebook:         {(sheet, name): {sheet, order, name, why, enter, format,
                                       allowed_raw, allowed, required, notes}}
    sheet_order:      {sheet: [(order, name), ...] sorted by order}
    controlled_vocab: {(sheet, name): [allowed, ...]}  (closed-vocab fields only)
    """
    import openpyxl
    cm = {**DEFAULT_COL_MAP, **(col_map or {})}
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    cb = wb[codebook_sheet]
    codebook, sheet_order, controlled_vocab = {}, {}, {}
    for r in range(2, cb.max_row + 1):
        sheet = cb.cell(r, cm["sheet"]).value
        order = cb.cell(r, cm["order"]).value
        name = cb.cell(r, cm["name"]).value
        if not sheet or not name:
            continue
        name = str(name).strip()
        allowed = split_allowed(cb.cell(r, cm["allowed"]).value)
        rec = {
            "sheet": sheet, "order": int(order), "name": name,
            "why": cb.cell(r, cm["why"]).value, "enter": cb.cell(r, cm["enter"]).value,
            "format": cb.cell(r, cm["format"]).value,
            "allowed_raw": cb.cell(r, cm["allowed"]).value,
            "allowed": allowed, "required": str(cb.cell(r, cm["required"]).value or "").strip(),
            "notes": cb.cell(r, cm["notes"]).value,
        }
        codebook[(sheet, name)] = rec
        sheet_order.setdefault(sheet, []).append((int(order), name))
        if allowed and allowed != [FREE_TEXT_SENTINEL]:
            controlled_vocab[(sheet, name)] = allowed
    for s in sheet_order:
        sheet_order[s].sort()
    return codebook, sheet_order, controlled_vocab


def controlled_vocab(codebook: dict) -> dict:
    """Extract {(sheet, name): [allowed,...]} for closed-vocabulary fields only."""
    out = {}
    for key, rec in codebook.items():
        allowed = rec.get("allowed") or []
        if allowed and allowed != [FREE_TEXT_SENTINEL]:
            out[key] = allowed
    return out


def verify_headers(workbook_path: str, sheet_order: dict, data_sheets) -> tuple:
    """Assert live sheet headers (row 1) match the Codebook order.

    Returns (ok: bool, problems: [(sheet, expected, actual), ...]).
    """
    import openpyxl
    wb = openpyxl.load_workbook(workbook_path, data_only=False)
    problems = []
    for sh in data_sheets:
        ws = wb[sh]
        expected = [name for _, name in sheet_order[sh]]
        actual = [str(ws.cell(1, c).value).strip() if ws.cell(1, c).value is not None else None
                  for c in range(1, len(expected) + 1)]
        if actual != expected:
            problems.append((sh, expected, actual))
    return (len(problems) == 0), problems
