"""Tests for services/extraction_excel_service.py — the coded extraction workbook."""
import io
import json
import uuid

import pytest
from openpyxl import load_workbook

from database import reset_db, init_db, get_connection
from config import settings
from services.db_backend import SQLiteBackend


@pytest.fixture(autouse=True)
def fresh_db(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "db_path", str(tmp_path / "test_xlsx.db"))
    monkeypatch.setattr(settings, "upload_root", str(tmp_path / "uploads"))
    monkeypatch.setattr(settings, "export_root", str(tmp_path / "exports"))
    reset_db()
    init_db()
    yield


def _db():
    return SQLiteBackend(get_connection())


def _setup():
    """Project with 2 sections x 2 variables, 3 studies (2 included, 1 excluded),
    and extractions covering DS / CALC / -99 / -77 / -88 / conflict."""
    db = _db()
    pid = str(uuid.uuid4())
    db.insert("projects", {"project_id": pid, "name": "Excel Review"})

    var_ids = {}
    specs = [
        ("Study_Metadata", "sample_size", "int"),
        ("Study_Metadata", "design", "enum"),
        ("Outcomes", "mean_age", "float"),
        ("Outcomes", "followup_months", "float"),
    ]
    for i, (section, name, ftype) in enumerate(specs):
        vid = str(uuid.uuid4())
        var_ids[name] = vid
        db.insert("variables", {
            "variable_id": vid, "project_id": pid, "section": section, "name": name,
            "field_type": ftype, "allowed_values": "", "description": "",
            "extraction_prompt": f"Extract {name} from the study.",
            "order_index": i + 1, "is_base": 0,
        })

    study_ids = {}
    for label, status in [("Included A", "included"), ("Included B", "included"),
                          ("Excluded C", "excluded")]:
        sid = str(uuid.uuid4())
        study_ids[label] = sid
        db.insert("studies", {
            "study_id": sid, "project_id": pid, "title": label,
            "abstract": "x", "screening_status": status,
        })

    def add_extraction(sid, vname, value, tag="", missing="", calc="", quote="", vstatus=""):
        db.insert("extractions", {
            "extraction_id": str(uuid.uuid4()),
            "study_id": sid, "variable_id": var_ids[vname],
            "value": value, "confidence": 0.9, "source_page": 3,
            "source_text": quote, "extracted_by": "coded_auto",
            "confidence_tag": tag, "missing_code": missing,
            "calc_note": calc, "extraction_prompt": "p",
            "verification_status": vstatus,
        })

    a, b = study_ids["Included A"], study_ids["Included B"]
    add_extraction(a, "sample_size", "200", tag="DS", quote="We enrolled 200 patients.")
    add_extraction(a, "design", "RCT", tag="DS", quote="randomized controlled trial")
    add_extraction(a, "mean_age", "-99", missing="-99")
    add_extraction(a, "followup_months", "24", tag="CALC", calc="2 yrs x 12 = 24 mo",
                   quote="followed for 2 years")
    add_extraction(b, "sample_size", "-77", missing="-77")
    add_extraction(b, "design", "-88", missing="-88")
    add_extraction(b, "mean_age", "55", tag="INF", vstatus="conflict")
    add_extraction(b, "followup_months", "12", tag="DS", quote="12-month follow-up")
    db.commit()
    return pid, study_ids, var_ids


def _wb(project_id):
    from services.extraction_excel_service import build_coded_extraction_workbook
    blob = build_coded_extraction_workbook(project_id, _db())
    return load_workbook(io.BytesIO(blob))


class TestWorkbookStructure:
    def test_one_sheet_per_section_plus_prompts_and_legend(self):
        pid, _, _ = _setup()
        wb = _wb(pid)
        assert "Study_Metadata" in wb.sheetnames
        assert "Outcomes" in wb.sheetnames
        assert "Prompts" in wb.sheetnames
        assert "Extraction_Legend" in wb.sheetnames

    def test_included_studies_only(self):
        pid, study_ids, _ = _setup()
        wb = _wb(pid)
        ws = wb["Study_Metadata"]
        ids = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        assert len(ids) == 2
        assert study_ids["Included A"][:8] in ids
        assert study_ids["Included B"][:8] in ids
        assert study_ids["Excluded C"][:8] not in ids

    def test_variables_as_columns(self):
        pid, _, _ = _setup()
        wb = _wb(pid)
        headers = [c.value for c in wb["Study_Metadata"][1]]
        assert headers[:2] == ["StudyID", "Study"]
        assert "sample_size" in headers
        assert "design" in headers
        assert headers[-1] == "Notes"


class TestMissingCodeColors:
    def _cell_fill(self, ws, row, col_name):
        headers = [c.value for c in ws[1]]
        col = headers.index(col_name) + 1
        return ws.cell(row=row, column=col).fill.fgColor.rgb

    def test_minus_99_yellow(self):
        pid, _, _ = _setup()
        ws = _wb(pid)["Outcomes"]
        assert self._cell_fill(ws, 2, "mean_age") == "00FFD700"

    def test_minus_77_grey_and_minus_88_orange(self):
        pid, _, _ = _setup()
        ws = _wb(pid)["Study_Metadata"]
        assert self._cell_fill(ws, 3, "sample_size") == "00D3D3D3"
        assert self._cell_fill(ws, 3, "design") == "00FFA500"

    def test_conflict_light_red(self):
        pid, _, _ = _setup()
        ws = _wb(pid)["Outcomes"]
        assert self._cell_fill(ws, 3, "mean_age") == "00FF6666"

    def test_normal_cells_uncolored(self):
        pid, _, _ = _setup()
        ws = _wb(pid)["Study_Metadata"]
        fill = self._cell_fill(ws, 2, "sample_size")
        assert fill in ("00000000", None) or fill == "FFFFFFFF"


class TestNotesAndLegend:
    def test_notes_column_has_tags_and_citations(self):
        pid, _, _ = _setup()
        ws = _wb(pid)["Outcomes"]
        headers = [c.value for c in ws[1]]
        notes_col = headers.index("Notes") + 1
        notes_a = ws.cell(row=2, column=notes_col).value
        assert "[CALC] followup_months: 2 yrs x 12 = 24 mo" in notes_a
        assert '[PDF-p.3: "followed for 2 years"]' in notes_a

    def test_legend_sheet_content(self):
        pid, _, _ = _setup()
        ws = _wb(pid)["Extraction_Legend"]
        codes = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        for expected in ["-99", "-77", "-88", "[DS]", "[INF]", "[CALC]", "[SUP]", "[FIG]"]:
            assert expected in codes

    def test_prompts_sheet_shows_extraction_prompts(self):
        pid, _, _ = _setup()
        ws = _wb(pid)["Prompts"]
        rows = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value
                for r in range(2, ws.max_row + 1)}
        assert rows["sample_size"] == "Extract sample_size from the study."
        assert "followup_months" in rows

    def test_workbook_round_trips_through_openpyxl(self):
        pid, _, _ = _setup()
        wb = _wb(pid)
        ws = wb["Study_Metadata"]
        headers = [c.value for c in ws[1]]
        col = headers.index("sample_size") + 1
        assert ws.cell(row=2, column=col).value == "200"
