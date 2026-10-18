"""Mixed-tool classification, grouping, profile, and ZIP export tests."""
import io
import json
import uuid
import zipfile

import pytest
import openpyxl

from config import settings
from database import get_connection, init_db, reset_db
import services.db_backend as db_mod
from services import rob_service
from services.db_backend import get_db_backend
from services.reporting_profile_service import assess_profile, persist_profile
from services.rob_classification_service import classify_study, persist_classification
from services.rob_export_service import export_project_package
from services.storage_backend import get_storage


@pytest.fixture(autouse=True)
def fresh_db(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "db_path", str(tmp_path / "test.db"))
    monkeypatch.setattr(settings, "upload_root", str(tmp_path / "uploads"))
    monkeypatch.setattr(settings, "export_root", str(tmp_path / "exports"))
    db_mod._backend_instance = None
    reset_db()
    init_db()
    yield
    reset_db()
    db_mod._backend_instance = None


def _project_with_studies():
    project_id = str(uuid.uuid4())
    studies = [
        {
            "study_id": str(uuid.uuid4()),
            "title": "Randomized treatment trial",
            "study_design": "randomized controlled trial",
            "abstract": "Participants underwent random allocation to treatment.",
        },
        {
            "study_id": str(uuid.uuid4()),
            "title": "A rare arterial case",
            "study_design": "case report",
            "abstract": "This case report describes a single patient.",
        },
    ]
    conn = get_connection()
    conn.execute("INSERT INTO projects (project_id, name) VALUES (?, ?)", (project_id, "Mixed"))
    for study in studies:
        conn.execute(
            """INSERT INTO studies
               (study_id, project_id, title, study_design, abstract, screening_status)
               VALUES (?, ?, ?, ?, ?, 'included')""",
            (study["study_id"], project_id, study["title"], study["study_design"], study["abstract"]),
        )
    conn.commit()
    conn.close()
    return project_id, studies


def test_ambiguous_classification_stays_manual():
    result = classify_study(
        {"title": "Narrative overview", "study_design": "", "abstract": "General discussion."},
        None,
        threshold=0.7,
    )
    assert result["tool"] == "manual"
    assert result["status"] == "manual_review"


def test_case_report_classification_selects_jbi():
    result = classify_study(
        {"title": "A case of disease", "study_design": "case report", "abstract": "single patient"},
        None,
    )
    assert result["tool"] == "jbi_case_report"
    assert result["status"] == "assigned"


def test_summary_groups_tools_and_exposes_manual_queue():
    project_id, studies = _project_with_studies()
    db = get_db_backend()
    rob_service.create_assessment(studies[0]["study_id"], project_id, "rob2", db=db)
    rob_service.create_assessment(studies[1]["study_id"], project_id, "jbi_case_report", db=db)
    persist_classification(db, studies[0]["study_id"], project_id, {
        "study_design": studies[0]["study_design"],
        "tool": "manual",
        "confidence": 0.4,
        "alternatives": ["rob2"],
        "rationale": "Conflicting design evidence.",
        "evidence": [],
        "status": "manual_review",
        "model_name": "test",
    })
    summary = rob_service.get_project_summary(project_id, db=db)
    assert summary["tool"] == "mixed"
    assert set(summary["tool_groups"]) == {"rob2", "jbi_case_report"}
    assert len(summary["manual_review"]) == 1


def test_profile_and_zip_package_have_expected_files():
    project_id, studies = _project_with_studies()
    db = get_db_backend()
    for study, tool in zip(studies, ("rob2", "jbi_case_report")):
        rob_service.create_assessment(study["study_id"], project_id, tool, db=db)
        classification = classify_study(study, None)
        persist_classification(db, study["study_id"], project_id, classification)
        items = assess_profile(study, {
            "full_text": (
                "MRI acquisition parameters included repetition time and echo time. "
                "The preprocessing pipeline used motion correction. CVR region of interest was defined."
            ),
            "pages": [],
        })
        persist_profile(db, project_id, study["study_id"], items)

    key = export_project_package(project_id, db)
    payload = get_storage("export").get_bytes(key)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = set(archive.namelist())
        assert "risk_of_bias.xlsx" in names
        assert "reporting_completeness.xlsx" in names
        assert "manifest.json" in names
        assert any(name.endswith(".svg") and "rob2" in name for name in names)
        assert any(name.endswith(".png") and "jbi_case_report" in name for name in names)
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["tool_groups"] == {"rob2": 1, "jbi_case_report": 1}

        workbook = openpyxl.load_workbook(io.BytesIO(archive.read("risk_of_bias.xlsx")))
        assert "Study Classification" in workbook.sheetnames
        assert "Manual Review" in workbook.sheetnames
        assert any(name.endswith("Summary") and "RoB 2" in name for name in workbook.sheetnames)
