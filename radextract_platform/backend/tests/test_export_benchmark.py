"""Benchmark test measuring performance of _export_csv and _export_xlsx in routers/export.py."""
import time
import io
import uuid
import pytest
from openpyxl import load_workbook

from database import reset_db, init_db, get_connection
from config import settings
from services.db_backend import SQLiteBackend
from routers.export import _export_csv, _export_xlsx


@pytest.fixture(autouse=True)
def benchmark_db(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "db_path", str(tmp_path / "test_benchmark.db"))
    monkeypatch.setattr(settings, "upload_root", str(tmp_path / "uploads"))
    monkeypatch.setattr(settings, "export_root", str(tmp_path / "exports"))
    reset_db()
    init_db()
    yield


def _db():
    return SQLiteBackend(get_connection())


def _setup_benchmark_data(num_studies=300, num_variables=15):
    db = _db()
    pid = f"proj_{uuid.uuid4().hex[:8]}"
    db.insert("projects", {"project_id": pid, "name": "Benchmark Project"})

    var_ids = []
    for i in range(num_variables):
        vid = f"var_{i}_{uuid.uuid4().hex[:6]}"
        var_ids.append(vid)
        db.insert("variables", {
            "variable_id": vid,
            "project_id": pid,
            "section": f"Section_{i % 3}",
            "name": f"Variable_{i}",
            "field_type": "free_text",
            "order_index": i,
        })

    for s_idx in range(num_studies):
        sid = f"study_{s_idx}_{uuid.uuid4().hex[:6]}"
        db.insert("studies", {
            "study_id": sid,
            "project_id": pid,
            "title": f"Study Title {s_idx}",
            "authors": f"Author A, Author B {s_idx}",
            "publication_year": 2020 + (s_idx % 5),
            "journal": f"Journal of Radiology {s_idx % 4}",
            "doi": f"10.1000/journal.2020.{s_idx}",
            "screening_status": "included",
        })

        # Insert extractions for each variable
        for v_idx, vid in enumerate(var_ids):
            db.insert("extractions", {
                "extraction_id": f"ext_{s_idx}_{v_idx}",
                "study_id": sid,
                "variable_id": vid,
                "value": f"val_{s_idx}_{v_idx}",
                "confidence": 0.95,
            })

        # Insert RoB assessment for study
        rid = f"rob_{s_idx}"
        db.insert("rob_assessments", {
            "assessment_id": rid,
            "study_id": sid,
            "project_id": pid,
            "tool": "ROBINS-I",
            "overall_judgment": "low_risk" if s_idx % 2 == 0 else "high_risk",
        })

    db.commit()
    return pid, num_studies, num_variables


def test_export_performance_benchmark():
    import asyncio

    pid, num_studies, num_variables = _setup_benchmark_data(num_studies=300, num_variables=15)
    db = _db()

    async def run_test():
        # Measure CSV export baseline time
        t0 = time.perf_counter()
        res_csv = await _export_csv(pid, "20260101_120000", include_rob=True, db=db)
        csv_time = time.perf_counter() - t0

        # Measure XLSX export baseline time
        t0 = time.perf_counter()
        res_xlsx = await _export_xlsx(pid, "20260101_120000", include_rob=True, db=db)
        xlsx_time = time.perf_counter() - t0

        return res_csv, res_xlsx, csv_time, xlsx_time

    res_csv, res_xlsx, csv_time, xlsx_time = asyncio.run(run_test())

    print(f"\n[BENCHMARK] CSV Export ({num_studies} studies, {num_variables} vars): {csv_time:.4f} seconds")
    print(f"[BENCHMARK] XLSX Export ({num_studies} studies, {num_variables} vars): {xlsx_time:.4f} seconds")

    # Basic assertions
    assert res_csv.status_code == 200
    assert len(res_csv.body) > 0
    assert res_xlsx.status_code == 200
    assert len(res_xlsx.body) > 0

    wb = load_workbook(io.BytesIO(res_xlsx.body))
    ws = wb["Extraction Data"]
    assert ws.max_row == num_studies + 1  # 1 header row + num_studies
