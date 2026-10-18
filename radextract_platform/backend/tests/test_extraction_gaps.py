"""Unit and integration tests for Extraction Quality Gaps (Part B).

Tests:
1. Confidence Taxonomy Tags ([DS], [INF], [CALC], [SUP], [FIG], [PDF-p.N: "quote"])
2. 7-Step Exhaustive Search & Mathematical Derivations
3. Targeted Recheck of Missing Data
4. Universal Gap Rules G1 to G5 (Fixed-horizon, comparator heterogeneity, substudy, immortal time, effect estimate resolution)
5. Scanned PDF Detection and Missing Data Codes (-99, -77, -88)
6. Missing Data Summary Aggregation
7. Excel Export Color Highlighting and Extraction_Legend sheet
"""

import io
import pytest
import sqlite3
import openpyxl

from services.exhaustive_search import (
    exhaustive_extract,
    recheck_missing,
    COMMON_ABBREVIATIONS,
)
from services.gap_rules import (
    apply_g1_fixed_horizon,
    apply_g2_comparator_heterogeneity,
    apply_g3_subgroup_detection,
    apply_g4_immortal_time_bias,
    apply_g5_effect_conflict,
)
from services.auto_extract_service import _is_scanned_pdf
from services.db_backend import SQLiteBackend


@pytest.fixture
def db_conn():
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    backend = SQLiteBackend(conn=conn)
    backend.execute("""
        CREATE TABLE projects (
            project_id TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            status TEXT DEFAULT 'active'
        )
    """)
    backend.execute("""
        CREATE TABLE variables (
            variable_id TEXT PRIMARY KEY,
            project_id TEXT NOT NULL,
            name TEXT NOT NULL,
            field_type TEXT DEFAULT 'text',
            section TEXT DEFAULT 'General',
            order_index INTEGER DEFAULT 0
        )
    """)
    backend.execute("""
        CREATE TABLE studies (
            study_id TEXT PRIMARY KEY,
            project_id TEXT NOT NULL,
            title TEXT NOT NULL,
            pdf_path TEXT,
            pdf_status TEXT DEFAULT 'ready'
        )
    """)
    backend.execute("""
        CREATE TABLE extractions (
            extraction_id TEXT PRIMARY KEY,
            study_id TEXT NOT NULL,
            variable_id TEXT NOT NULL,
            value TEXT,
            confidence REAL,
            confidence_tag TEXT,
            source_page INTEGER,
            source_bbox TEXT,
            source_text TEXT,
            extracted_by TEXT
        )
    """)
    backend.commit()
    return backend


class TestMathematicalDerivations:
    def test_percentage_female_from_male(self):
        existing = {"pct_male": "40"}
        res = exhaustive_extract("pct_female", "Some text", existing_values=existing)
        assert res is not None
        assert res["value"] == "60.0"
        assert res["confidence_tag"] == "[CALC]"

    def test_total_n_from_arms(self):
        existing = {"n_arm1": "120", "n_arm2": "130"}
        res = exhaustive_extract("total_n", "Some text", existing_values=existing)
        assert res is not None
        assert res["value"] == "250"
        assert res["confidence_tag"] == "[CALC]"

    def test_standard_deviation_from_iqr(self):
        existing = {"iqr": "13.5"}
        res = exhaustive_extract("sd", "Some text", existing_values=existing)
        assert res is not None
        assert res["value"] == "10.0"
        assert res["confidence_tag"] == "[CALC]"


class TestSynonymAndAbbreviationResolution:
    def test_abbreviation_lookup(self):
        text = "The primary endpoint showed hazard ratio of 0.72."
        res = exhaustive_extract("hr", text)
        assert res is not None
        assert "0.72" in res["value"]

    def test_synonym_lookup(self):
        text = "The area under the curve was 0.88."
        res = exhaustive_extract("auc", text)
        assert res is not None
        assert "0.88" in res["value"]


class TestExhaustiveSearchPipeline:
    def test_fallback_through_search_locations(self):
        abstract = "Abstract: A total of 350 patients were enrolled."
        full_text = "Results: Primary outcome metric was overall survival of 14.2 months."
        existing = {"sample_size": "350"}

        # Step 1 structured
        res1 = exhaustive_extract("sample_size", full_text, existing_values=existing)
        assert res1["value"] == "350"
        assert res1["confidence_tag"] == "[DS]"

        # Step 2 abstract
        res2 = exhaustive_extract("enrolled", abstract + full_text)
        assert res2 is not None

        # Step 4 results
        res3 = exhaustive_extract("survival", full_text)
        assert res3 is not None

    def test_recheck_missing_pass(self):
        extractions = [
            {
                "name": "followup",
                "value": "-99",
                "confidence": 0.0,
            }
        ]
        full_text = "Patients were followed for a median duration of 24.5 months."

        updated = recheck_missing(extractions, full_text)
        assert len(updated) == 1
        assert updated[0]["value"] != "-99"
        assert "24.5" in updated[0]["value"]
        assert updated[0]["confidence_tag"] == "[RECHECK]"


class TestGapRulesG1toG5:
    def test_g1_fixed_horizon_followup(self):
        text = "All patients were followed for 2.0 years — no reported median."
        extractions = [{"variable_name": "median_followup_months", "value": "-99"}]
        res = apply_g1_fixed_horizon(extractions, text)
        assert len(res) == 1
        assert res[0]["value"] == "24"
        assert res[0]["confidence_tag"] == "[CALC]"

    def test_g2_comparator_heterogeneity(self):
        text = "The control arm consisted of patients receiving warfarin or DOAC."
        extractions = [{"variable_name": "comparator_type", "value": ""}]
        res = apply_g2_comparator_heterogeneity(extractions, text)
        assert len(res) == 1
        assert res[0]["value"] == "mixed"
        assert res[0]["confidence_tag"] == "[INF]"

    def test_g3_subgroup_detection(self):
        text = "Patients enrolled in this post-hoc analysis of KEYNOTE-001 trial were studied."
        extractions = [{"variable_name": "is_subgroup", "value": ""}]
        res = apply_g3_subgroup_detection(extractions, text)
        assert len(res) == 1
        assert res[0]["value"] == "Y"
        assert res[0]["confidence_tag"] == "[INF]"

    def test_g4_immortal_time_bias(self):
        text = "This observational study evaluated patients where follow-up started after procedure."
        extractions = [{"variable_name": "risk_of_bias", "value": "", "source_text": ""}]
        res = apply_g4_immortal_time_bias(extractions, text)
        assert len(res) == 1
        assert "IMMORTAL TIME BIAS" in res[0]["source_text"]

    def test_g5_effect_conflict_resolution(self):
        text = "Hazard ratio was reported as HR = 0.65 in abstract and HR = 0.80 in subgroup."
        extractions = [{"variable_name": "hr", "value": "0.65"}]
        res = apply_g5_effect_conflict(extractions, text)
        assert len(res) == 1
        assert res[0]["confidence_tag"] == "[CONFLICT]"
        assert res[0]["value"] == "0.8"  # closer to 1.0 (conservative)


class TestScannedPdfDetection:
    def test_scanned_pdf_identification(self):
        assert _is_scanned_pdf({"pages": [{"page_number": 1}], "full_text": ""}) is True
        assert _is_scanned_pdf({"pages": [{"page_number": 1}], "full_text": "Short header only"}) is True  # < 100 chars
        assert _is_scanned_pdf({"pages": [{"page_number": 1}], "full_text": "This is a full article with plenty of text " * 20}) is False


class TestMissingDataSummaryAndExportStyling:
    def test_missing_data_summary_calculation(self, db_conn):
        from routers.extraction import get_missing_data_summary

        # Insert project and variables
        db_conn.execute("INSERT INTO projects (project_id, name) VALUES ('p1', 'Test')")
        db_conn.execute("INSERT INTO variables (variable_id, project_id, name) VALUES ('v1', 'p1', 'Age')")
        db_conn.execute("INSERT INTO variables (variable_id, project_id, name) VALUES ('v2', 'p1', 'Dose')")
        db_conn.execute("INSERT INTO studies (study_id, project_id, title) VALUES ('s1', 'p1', 'Study 1')")
        db_conn.execute("INSERT INTO studies (study_id, project_id, title) VALUES ('s2', 'p1', 'Study 2')")

        # Extractions: one valid, one -99, one -77, one -88
        db_conn.execute("INSERT INTO extractions (extraction_id, study_id, variable_id, value) VALUES ('e1', 's1', 'v1', '65')")
        db_conn.execute("INSERT INTO extractions (extraction_id, study_id, variable_id, value) VALUES ('e2', 's1', 'v2', '-99')")
        db_conn.execute("INSERT INTO extractions (extraction_id, study_id, variable_id, value) VALUES ('e3', 's2', 'v1', '-77')")
        db_conn.execute("INSERT INTO extractions (extraction_id, study_id, variable_id, value) VALUES ('e4', 's2', 'v2', '-88')")
        db_conn.commit()

        import asyncio
        summary = asyncio.run(get_missing_data_summary(project_id="p1", user=None, conn=db_conn))

        assert summary.total_extractions == 4
        assert summary.not_reported == 1
        assert summary.not_applicable == 1
        assert summary.unable_to_extract == 1
        assert len(summary.per_variable) == 2

    def test_xlsx_export_pattern_fill_and_legend(self, db_conn):
        from routers.export import _export_xlsx

        db_conn.execute("INSERT INTO projects (project_id, name) VALUES ('p1', 'Test')")
        db_conn.execute("INSERT INTO variables (variable_id, project_id, name) VALUES ('v1', 'p1', 'Age')")
        db_conn.execute("INSERT INTO studies (study_id, project_id, title) VALUES ('s1', 'p1', 'Study 1')")
        db_conn.execute("INSERT INTO extractions (extraction_id, study_id, variable_id, value) VALUES ('e1', 's1', 'v1', '-99')")
        db_conn.commit()

        import asyncio
        response = asyncio.run(_export_xlsx("p1", "20260829_000000", include_rob=False, db=db_conn))

        assert response.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(response.body))
        
        # Verify worksheets
        assert "Extraction Data" in wb.sheetnames
        assert "Extraction_Legend" in wb.sheetnames

        # Check legend contents
        ws_legend = wb["Extraction_Legend"]
        assert ws_legend.cell(row=1, column=1).value == "Code"
        assert ws_legend.cell(row=2, column=1).value == "-99"
        assert ws_legend.cell(row=3, column=1).value == "-77"
        assert ws_legend.cell(row=4, column=1).value == "-88"

        # Check color fill on -99
        ws_data = wb["Extraction Data"]
        found_fill = False
        for row in ws_data.iter_rows():
            for cell in row:
                if str(cell.value) == "-99":
                    assert cell.fill.start_color.rgb in ("00FFD700", "FFD700")
                    found_fill = True
        assert found_fill is True
