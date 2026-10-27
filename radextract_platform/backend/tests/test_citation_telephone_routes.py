"""Citation routes must never turn bibliographies or offline models into evidence."""
import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from routers import citation_telephone as routes


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(routes.router)
    return TestClient(app)


def test_upload_extracts_all_claims_without_evaluating_references(client, monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Bibliographic entries must never be evaluated as passages")
    monkeypatch.setattr(routes, "get_laya_engine", forbidden)
    monkeypatch.setattr(routes, "evaluate_passage_with_laya", forbidden)
    text = "\n".join(f"This is the claim for study number {n} [{n}]." for n in range(1, 18))
    text += "\nReferences\n" + "\n".join(f"{n}. Author. Source title {n}." for n in range(1, 18))
    response = client.post("/api/citation-telephone/upload-manuscript", files={"file": ("draft.TXT", text.encode(), "text/plain")})
    assert response.status_code == 200
    result = response.json()
    assert result["totalClaimsFound"] == len(result["claims"]) == 17
    assert result["totalReferencesFound"] == 17
    assert result["verificationStatus"] == "source_passages_required"
    assert all("verdict" not in claim for claim in result["claims"])


@pytest.mark.parametrize("filename,content,status", [("draft.pdf", b"%PDF-1.4", 415), ("draft.docx", b"broken", 400), ("draft.txt", b"\xff", 400), ("draft.txt", b" ", 400)])
def test_invalid_upload_is_not_reported_as_success(client, filename, content, status):
    response = client.post("/api/citation-telephone/upload-manuscript", files={"file": (filename, content)})
    assert response.status_code == status


def test_docx_upload_parses_real_paragraphs(client):
    from docx import Document
    doc = Document()
    doc.add_paragraph("This intervention improved recovery [1].")
    doc.add_paragraph("References")
    doc.add_paragraph("1. Smith. Recovery study.")
    content = io.BytesIO()
    doc.save(content)
    response = client.post("/api/citation-telephone/upload-manuscript", files={"file": ("draft.docx", content.getvalue())})
    assert response.status_code == 200
    assert response.json()["totalClaimsFound"] == 1


def test_paper_id_validation_prevents_ssrf_and_injection(client):
    # Empty ID
    res = client.get("/api/citation-telephone/paper?id= ")
    assert res.status_code == 400

    # Query injection / parameter manipulation
    res = client.get("/api/citation-telephone/paper?id=10.1000/182%26extra%3D1")
    assert res.status_code == 400

    # Space or illegal characters in identifier
    res = client.get("/api/citation-telephone/paper?id=PMC 123456")
    assert res.status_code == 400

    # Path traversal sequence
    res = client.get("/api/citation-telephone/paper?id=..%2F..%2Fetc%2Fpasswd")
    assert res.status_code == 400

    # Overly long string
    long_id = "A" * 201
    res = client.get(f"/api/citation-telephone/paper?id={long_id}")
    assert res.status_code == 400


def test_empty_passage_is_rejected(client):
    response = client.post("/api/citation-telephone/evaluate-claim", json={"citing_sentence": "A claim", "source_passage": " "})
    assert response.status_code == 400


def test_offline_model_produces_no_verdict(client, monkeypatch):
    monkeypatch.setattr(routes, "get_laya_engine", lambda: None)
    response = client.post("/api/citation-telephone/evaluate-claim", json={"citing_sentence": "A claim", "source_passage": "Actual source passage"})
    assert response.status_code == 503
    assert "verdict" not in response.json()


def test_inference_fallback_is_rejected(client, monkeypatch):
    monkeypatch.setattr(routes, "get_laya_engine", object)
    monkeypatch.setattr(routes, "evaluate_passage_with_laya", lambda **kwargs: {"engine": "heuristic_fallback", "verdict": "same"})
    response = client.post("/api/citation-telephone/evaluate-claim", json={"citing_sentence": "A claim", "source_passage": "Actual source passage"})
    assert response.status_code == 503


def test_evaluation_uses_supplied_passage(client, monkeypatch):
    engine = object()
    monkeypatch.setattr(routes, "get_laya_engine", lambda: engine)
    def evaluate(**kwargs):
        assert kwargs == {"citing_sentence": "A claim", "source_title": "Source title", "source_passage": "Actual source passage", "laya": engine}
        return {"engine": "local_laya_system_1", "verdict": "different"}
    monkeypatch.setattr(routes, "evaluate_passage_with_laya", evaluate)
    response = client.post("/api/citation-telephone/evaluate-claim", json={"citing_sentence": "A claim", "source_title": "Source title", "source_passage": "Actual source passage"})
    assert response.status_code == 200
    assert response.json()["verdict"] == "different"
