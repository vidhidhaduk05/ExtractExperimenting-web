"""Tests for PRISMA PNG/DOCX export (services/prisma_export_service.py + endpoints)."""
import io
import pytest
from fastapi.testclient import TestClient
from main import app
from database import reset_db, init_db
from config import settings


@pytest.fixture(autouse=True)
def fresh_db(tmp_path, monkeypatch):
    db_path = str(tmp_path / "test_api.db")
    monkeypatch.setattr(settings, "db_path", db_path)
    monkeypatch.setattr(settings, "upload_root", str(tmp_path / "uploads"))
    monkeypatch.setattr(settings, "export_root", str(tmp_path / "exports"))
    reset_db()
    init_db()
    yield


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def project_with_studies(client):
    """Project with a mix of pending / TA-excluded / FT-excluded / included studies."""
    r = client.post("/api/projects", json={"name": "PRISMA Export Review"})
    pid = r.json()["project_id"]

    def add(title, stage, status, reason=""):
        r = client.post(f"/api/projects/{pid}/studies", json={
            "project_id": pid, "title": title,
        })
        sid = r.json()["study_id"]
        if stage or status:
            client.put(f"/api/studies/{sid}", json={
                "screening_stage": stage, "screening_status": status,
                "screening_reason": reason,
            })
        return sid

    add("Pending study", "ta", "pending")
    add("TA excluded wrong population", "ta", "excluded", "Wrong population")
    add("TA excluded wrong design", "ta", "excluded", "Wrong design")
    add("FT excluded no outcome", "ft", "excluded", "Outcome not reported")
    add("Included study A", "ft", "included")
    add("Included study B", "ft", "included")
    return pid


def _service_db():
    """Raw connection wrapped in DBBackend for direct service calls."""
    from database import get_connection
    from services.db_backend import SQLiteBackend
    return SQLiteBackend(get_connection())


class TestFlowData:
    def test_flow_data_counts(self, client, project_with_studies):
        from services.prisma_export_service import get_prisma_flow_data
        data = get_prisma_flow_data(project_with_studies, _service_db())
        assert data["identification"]["total"] == 6
        assert data["screening"]["records_screened"] == 6
        assert data["screening"]["records_excluded"] == 2
        assert data["eligibility"]["full_text_assessed"] == 3
        assert data["eligibility"]["full_text_excluded"] == 1
        assert data["included"]["studies_in_review"] == 2
        assert data["pending"] == 1
        reasons = {r["reason"]: r["count"] for r in data["screening"]["excluded_with_reasons"]}
        assert reasons == {"Wrong population": 1, "Wrong design": 1}

    def test_json_endpoint_matches_service(self, client, project_with_studies):
        from services.prisma_export_service import get_prisma_flow_data
        r = client.get(f"/api/projects/{project_with_studies}/prisma")
        assert r.status_code == 200
        assert r.json() == get_prisma_flow_data(project_with_studies, _service_db())


class TestPngExport:
    def test_png_endpoint_returns_image(self, client, project_with_studies):
        r = client.get(f"/api/projects/{project_with_studies}/prisma/png")
        assert r.status_code == 200
        assert r.headers["content-type"] == "image/png"
        # PNG magic bytes
        assert r.content[:8] == b"\x89PNG\r\n\x1a\n"
        assert len(r.content) > 1000
        assert "attachment" in r.headers.get("content-disposition", "")


class TestDocxExport:
    def test_docx_endpoint_returns_word_file(self, client, project_with_studies):
        r = client.get(f"/api/projects/{project_with_studies}/prisma/docx")
        assert r.status_code == 200
        assert r.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        # DOCX is a zip
        assert r.content[:2] == b"PK"

    def test_docx_contains_editable_tables(self, client, project_with_studies):
        from docx import Document
        r = client.get(f"/api/projects/{project_with_studies}/prisma/docx")
        doc = Document(io.BytesIO(r.content))
        # 6 tables: identification, screening, TA reasons, eligibility, FT reasons, included
        assert len(doc.tables) == 6
        # Identification table holds the total
        id_table = doc.tables[0]
        pairs = {row.cells[0].text: row.cells[1].text for row in id_table.rows}
        assert pairs["Total records identified"] == "6"
        # Included table
        incl_table = doc.tables[5]
        pairs = {row.cells[0].text: row.cells[1].text for row in incl_table.rows}
        assert pairs["Studies included in review"] == "2"
        assert pairs["Records pending screening"] == "1"

    def test_docx_exclusion_reasons_match_flow(self, client, project_with_studies):
        from docx import Document
        r = client.get(f"/api/projects/{project_with_studies}/prisma/docx")
        doc = Document(io.BytesIO(r.content))
        ta_table = doc.tables[2]
        reasons = {row.cells[0].text for row in ta_table.rows[1:]}
        assert reasons == {"Wrong population", "Wrong design"}
        ft_table = doc.tables[4]
        reasons = {row.cells[0].text for row in ft_table.rows[1:]}
        assert reasons == {"Outcome not reported"}

    def test_docx_counts_match_json_endpoint(self, client, project_with_studies):
        from docx import Document
        flow = client.get(f"/api/projects/{project_with_studies}/prisma").json()
        r = client.get(f"/api/projects/{project_with_studies}/prisma/docx")
        doc = Document(io.BytesIO(r.content))
        scr_table = doc.tables[1]
        pairs = {row.cells[0].text: row.cells[1].text for row in scr_table.rows}
        assert int(pairs["Records screened"]) == flow["screening"]["records_screened"]
        assert int(pairs["Records excluded"]) == flow["screening"]["records_excluded"]
