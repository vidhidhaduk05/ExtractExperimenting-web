"""Tests for Excel-style Extraction Workspace:
1. Authoritative variable ordering & reordering.
2. Protection against accidental deletion of populated variables.
3. Batch extraction grid retrieval and cell upsert.
4. Formula injection sanitization and dual Excel exports (.xlsx).
"""
import io
import uuid
import pytest
from database import init_sqlite_db, get_connection
from services.db_backend import SQLiteBackend
import openpyxl


@pytest.fixture
def test_db():
    conn = get_connection()
    init_sqlite_db(conn)
    backend = SQLiteBackend(conn)
    yield backend
    conn.close()


@pytest.mark.asyncio
async def test_variable_reorder_and_ordering(test_db):
    from routers.variables import list_variables, create_variable, reorder_variables
    from models import VariableCreate, VariableReorder

    project_id = f"proj_test_order_{uuid.uuid4().hex[:8]}"
    test_db.execute("INSERT INTO projects (project_id, name) VALUES (?, ?)", (project_id, "Test Project"))
    test_db.commit()

    # 1. Create variables in sequence: Study ID, Year, Author, Country, Sample size
    var_names = ["Study ID", "Year", "Author", "Country", "Sample size"]
    created_ids = []
    for idx, name in enumerate(var_names):
        res = await create_variable(
            project_id=project_id,
            body=VariableCreate(name=name, section="Demographics", order_index=idx),
            user=None,
            conn=test_db,
        )
        created_ids.append(res.variable_id)

    # Verify listing returns exact sequence
    listed = await list_variables(project_id, user=None, conn=test_db)
    assert [v["name"] for v in listed] == var_names

    # 2. Reorder Country before Author: Study ID, Year, Country, Author, Sample size
    reordered_ids = [created_ids[0], created_ids[1], created_ids[3], created_ids[2], created_ids[4]]
    reorder_res = await reorder_variables(
        project_id=project_id,
        body=VariableReorder(variable_ids=reordered_ids),
        user=None,
        conn=test_db,
    )
    assert [v["name"] for v in reorder_res] == ["Study ID", "Year", "Country", "Author", "Sample size"]

    # Re-fetch independently to verify database persistence
    persisted = await list_variables(project_id, user=None, conn=test_db)
    assert [v["name"] for v in persisted] == ["Study ID", "Year", "Country", "Author", "Sample size"]


@pytest.mark.asyncio
async def test_protected_variable_deletion(test_db):
    from routers.variables import create_variable, delete_variable
    from models import VariableCreate
    from fastapi import HTTPException

    project_id = f"proj_test_del_{uuid.uuid4().hex[:8]}"
    test_db.execute("INSERT INTO projects (project_id, name) VALUES (?, ?)", (project_id, "Delete Test Project"))
    test_db.commit()

    var_res = await create_variable(
        project_id=project_id,
        body=VariableCreate(name="Populated Var", section="General"),
        user=None,
        conn=test_db,
    )
    var_id = var_res.variable_id

    # Add populated extraction value
    study_id = f"study_del_{uuid.uuid4().hex[:8]}"
    test_db.execute(
        "INSERT INTO studies (study_id, project_id, title) VALUES (?, ?, ?)",
        (study_id, project_id, "Study 1"),
    )
    test_db.execute(
        "INSERT INTO extractions (extraction_id, study_id, variable_id, value) VALUES (?, ?, ?, ?)",
        (f"ext_del_{uuid.uuid4().hex[:8]}", study_id, var_id, "Critical Clinical Measurement"),
    )
    test_db.commit()

    # Attempt delete without force=True -> must raise 409 Conflict
    with pytest.raises(HTTPException) as exc_info:
        await delete_variable(var_id, force=False, user=None, conn=test_db)
    assert exc_info.value.status_code == 409
    assert exc_info.value.detail["requires_confirmation"] is True
    assert exc_info.value.detail["populated_count"] == 1

    # Now delete with force=True -> succeeds
    del_res = await delete_variable(var_id, force=True, user=None, conn=test_db)
    assert del_res["deleted"] is True
    assert del_res["deleted_extractions_count"] == 1


@pytest.mark.asyncio
async def test_extraction_grid_and_batch_upsert(test_db):
    from routers.extraction import get_extraction_grid, batch_update_extractions
    from models import ExtractionBatchUpsert, ExtractionBatchCell

    project_id = f"proj_test_grid_{uuid.uuid4().hex[:8]}"
    study_g1 = f"study_g1_{uuid.uuid4().hex[:8]}"
    study_g2 = f"study_g2_{uuid.uuid4().hex[:8]}"
    var_sample = f"var_sample_{uuid.uuid4().hex[:8]}"
    var_country = f"var_country_{uuid.uuid4().hex[:8]}"

    test_db.execute("INSERT INTO projects (project_id, name) VALUES (?, ?)", (project_id, "Grid Test Project"))
    test_db.execute("INSERT INTO studies (study_id, project_id, title, authors, publication_year, study_label) VALUES (?, ?, ?, ?, ?, ?)",
                    (study_g1, project_id, "Study Alpha", "Smith et al.", 2024, "Smith 2024"))
    test_db.execute("INSERT INTO studies (study_id, project_id, title, authors, publication_year, study_label) VALUES (?, ?, ?, ?, ?, ?)",
                    (study_g2, project_id, "Study Beta", "Jones et al.", 2023, "Jones 2023"))
    test_db.execute("INSERT INTO variables (variable_id, project_id, name, order_index, section, field_type) VALUES (?, ?, ?, ?, ?, ?)",
                    (var_sample, project_id, "Sample size", 1, "Characteristics", "numeric"))
    test_db.execute("INSERT INTO variables (variable_id, project_id, name, order_index, section, field_type) VALUES (?, ?, ?, ?, ?, ?)",
                    (var_country, project_id, "Country", 2, "Characteristics", "free_text"))
    test_db.commit()

    # 1. Batch upsert rectangular cells
    batch_data = ExtractionBatchUpsert(
        cells=[
            ExtractionBatchCell(study_id=study_g1, variable_id=var_sample, value="150"),
            ExtractionBatchCell(study_id=study_g1, variable_id=var_country, value="USA"),
            ExtractionBatchCell(study_id=study_g2, variable_id=var_sample, value="75"),
            ExtractionBatchCell(study_id=study_g2, variable_id=var_country, value="Canada"),
        ]
    )
    upsert_res = await batch_update_extractions(project_id, batch_data, user=None, conn=test_db)
    assert upsert_res["status"] == "success"
    assert upsert_res["created_count"] == 4

    # 2. Retrieve grid
    grid = await get_extraction_grid(project_id, user=None, conn=test_db)
    assert len(grid["studies"]) == 2
    assert len(grid["variables"]) == 2
    assert grid["extractions"][study_g1][var_sample]["value"] == "150"
    assert grid["extractions"][study_g1][var_country]["value"] == "USA"
    assert grid["extractions"][study_g2][var_sample]["value"] == "75"
    assert grid["extractions"][study_g2][var_country]["value"] == "Canada"


@pytest.mark.asyncio
async def test_excel_exports_workbook_and_template(test_db):
    from routers.export import export_blank_template, _export_xlsx
    from models import VariableCreate
    from routers.variables import create_variable

    project_id = f"proj_test_xlsx_{uuid.uuid4().hex[:8]}"
    study_e1 = f"study_e1_{uuid.uuid4().hex[:8]}"

    test_db.execute("INSERT INTO projects (project_id, name) VALUES (?, ?)", (project_id, "Export Project"))
    test_db.execute("INSERT INTO studies (study_id, project_id, title, authors, publication_year, study_label) VALUES (?, ?, ?, ?, ?, ?)",
                    (study_e1, project_id, "Vascular Study", "Doe et al.", 2024, "Doe 2024"))

    var_order = ["Year", "Author", "Country", "Sample size"]
    for idx, name in enumerate(var_order):
        await create_variable(
            project_id=project_id,
            body=VariableCreate(name=name, section="General", order_index=idx),
            user=None,
            conn=test_db,
        )

    # 1. Blank Template Export
    template_res = await export_blank_template(project_id, user=None, db=test_db)
    assert template_res.status_code == 200
    tpl_wb = openpyxl.load_workbook(io.BytesIO(template_res.body))
    tpl_ws = tpl_wb["Extraction Sheet"]
    # Row 1 headers must be Study ID followed by variables in saved order
    headers = [cell.value for cell in tpl_ws[1]]
    assert headers == ["Study ID", "Year", "Author", "Country", "Sample size"]
    assert "Variable Definitions" in tpl_wb.sheetnames

    # 2. Completed Workbook Export with formula injection check
    # Add a cell with formula injection syntax: =SUM(1,2)
    ext_id = f"ext_inj_{uuid.uuid4().hex[:8]}"
    test_db.execute(
        """INSERT INTO extractions
           (extraction_id, study_id, variable_id, value, source_text, source_page, confidence)
           SELECT ?, ?, variable_id, '=SUM(1,2)', 'Sample quote on page 4', 4, 0.98
           FROM variables WHERE project_id = ? AND name = 'Country'""",
        (ext_id, study_e1, project_id),
    )
    test_db.commit()

    completed_res = await _export_xlsx(project_id, "20261026", include_rob=False, db=test_db)
    comp_wb = openpyxl.load_workbook(io.BytesIO(completed_res.body))
    comp_ws = comp_wb["Extraction Data"]

    # Verify Main Sheet
    assert [cell.value for cell in comp_ws[1]] == ["Study ID", "Year", "Author", "Country", "Sample size"]
    assert comp_ws.cell(row=2, column=1).value == "Doe 2024"
    # Country column is column 4: formula injection was escaped with leading single quote
    country_val = comp_ws.cell(row=2, column=4).value
    assert country_val == "'=SUM(1,2)"

    # Verify Evidence Sheet
    assert "Evidence" in comp_wb.sheetnames
    ev_ws = comp_wb["Evidence"]
    ev_headers = [cell.value for cell in ev_ws[1]]
    assert ev_headers[:5] == ["Study ID", "Study Label", "Variable ID", "Variable Name", "Extracted Value"]
    assert ev_ws.cell(row=2, column=2).value == "Doe 2024"
    assert ev_ws.cell(row=2, column=6).value == "Sample quote on page 4"
    assert ev_ws.cell(row=2, column=8).value == 4
