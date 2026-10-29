from __future__ import annotations
import csv
import io
import os
import uuid
from datetime import datetime
from fastapi import APIRouter, HTTPException, Query, Depends
from fastapi.responses import FileResponse, Response

from database import get_db
from services.db_backend import DBBackend
from dependencies import get_current_user_optional, check_project_access
from config import settings
from services.storage_backend import get_storage
from services.reference_export import export_ris, export_bibtex, export_endnote_xml

router = APIRouter(prefix="/api", tags=["export"])


@router.get("/projects/{project_id}/export/template.xlsx")
async def export_template_xlsx_alias(
    project_id: str,
    user: dict = Depends(get_current_user_optional),
    db: DBBackend = Depends(get_db),
):
    """Export blank extraction .xlsx template with column headers in saved sequence and dropdown validation."""
    return await export_blank_template(project_id, user=user, db=db)


@router.get("/projects/{project_id}/export/template")
async def export_blank_template(
    project_id: str,
    user: dict = Depends(get_current_user_optional),
    db: DBBackend = Depends(get_db),
):
    """Generate blank extraction template (.xlsx) driven by project's saved variable list."""
    check_project_access(db, project_id, user)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    variables = db.query("variables", {"project_id": project_id}, order_by="order_index ASC, name ASC")
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extraction Sheet"

    # Determine headers: check if "Study ID" already exists as a variable
    has_study_id = any(v.get("name", "").strip().lower() in ("study id", "study_id") for v in variables)
    col_vars = []
    headers = []
    if not has_study_id:
        headers.append("Study ID")
    for v in variables:
        vname = v.get("name", "").strip()
        headers.append(vname)
        col_vars.append(v)

    ws.append(headers)

    # Style header row (Row 1)
    header_fill = PatternFill(start_color="3A0F34", end_color="3A0F34", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FBF9F4")
    thin_border = Border(
        left=Side(style='thin', color='D2C2CB'),
        right=Side(style='thin', color='D2C2CB'),
        top=Side(style='thin', color='D2C2CB'),
        bottom=Side(style='thin', color='D2C2CB'),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(18, len(str(cell.value or "")) + 4)

    # Add 5 blank sample rows with borders
    for row_idx in range(2, 7):
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.border = thin_border

    # Freeze header row and Study ID column
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}6"

    # Add dropdown validations for categorical variables
    start_var_offset = 1 if not has_study_id else 0
    for idx, v in enumerate(col_vars):
        allowed_str = v.get("allowed_values", "")
        if allowed_str and v.get("field_type") == "categorical":
            opts = [x.strip() for x in allowed_str.replace(";", ",").split(",") if x.strip()]
            if opts:
                formula_str = f'"{",".join(opts)}"'
                dv = DataValidation(type="list", formula1=formula_str, allow_blank=True)
                ws.add_data_validation(dv)
                col_letter = openpyxl.utils.get_column_letter(start_var_offset + idx + 1)
                dv.add(f"{col_letter}2:{col_letter}500")

    # Instructions Sheet
    ws_guide = wb.create_sheet(title="Variable Definitions")
    ws_guide.append(["Column #", "Variable Name", "Section", "Data Type", "Allowed Values", "Description / Extraction Prompt"])
    for c_idx in range(1, 7):
        ws_guide.cell(row=1, column=c_idx).font = Font(bold=True)
    for idx, v in enumerate(col_vars, start=1):
        ws_guide.append([
            idx,
            v.get("name", ""),
            v.get("section", "General"),
            v.get("field_type", "free_text"),
            v.get("allowed_values", ""),
            v.get("description", ""),
        ])

    out_stream = io.BytesIO()
    wb.save(out_stream)
    xlsx_bytes = out_stream.getvalue()
    filename = f"extraction_template_{project_id}.xlsx"
    return Response(
        content=xlsx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/projects/{project_id}/export/csv")
async def export_project_csv_alias(
    project_id: str,
    include_rob: bool = Query(True, description="Include RoB data in export"),
    user: dict = Depends(get_current_user_optional),
    db: DBBackend = Depends(get_db),
):
    """Alias for GET /projects/{id}/export?format=csv."""
    return await export_project(project_id, format="csv", include_rob=include_rob, user=user, db=db)


@router.get("/projects/{project_id}/export/xlsx")
async def export_project_xlsx_alias(
    project_id: str,
    include_rob: bool = Query(True, description="Include RoB data in export"),
    user: dict = Depends(get_current_user_optional),
    db: DBBackend = Depends(get_db),
):
    """Alias for GET /projects/{id}/export?format=xlsx."""
    return await export_project(project_id, format="xlsx", include_rob=include_rob, user=user, db=db)


@router.get("/projects/{project_id}/export")
async def export_project(
    project_id: str,
    format: str = Query("csv", pattern="^(csv|xlsx)$"),
    include_rob: bool = Query(True, description="Include RoB data in export"),
    user: dict = Depends(get_current_user_optional),
    db: DBBackend = Depends(get_db),
):
    """Export all project data (extractions + RoB) as CSV or Excel."""
    check_project_access(db, project_id, user)
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")

    if format == "xlsx":
        return await _export_xlsx(project_id, timestamp, include_rob, db)
    else:
        return await _export_csv(project_id, timestamp, include_rob, db)


def _sanitize_cell_value(val):
    """Format value and guard against CSV/Excel formula injection."""
    if val is None:
        return ""
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    if not s:
        return ""
    # Try integer/float conversion
    try:
        if s.isdigit() or (s.startswith("-") and s[1:].isdigit()):
            return int(s)
        f_val = float(s)
        return f_val
    except ValueError:
        pass
    # Guard against formula execution in spreadsheet viewers (=, +, -, @)
    if s.startswith(("=", "+", "-", "@")):
        return f"'{s}"
    return s


def _fetch_export_data(project_id: str, include_rob: bool, db: DBBackend, study_id: Optional[str] = None):
    proj = db.get("projects", project_id)
    if not proj:
        raise HTTPException(status_code=404, detail="Project not found")

    variables = db.query("variables", {"project_id": project_id}, order_by="order_index ASC, name ASC")
    if study_id:
        single_study = db.get("studies", study_id)
        if not single_study or single_study.get("project_id") != project_id:
            raise HTTPException(status_code=404, detail="Study not found in project")
        studies = [single_study]
    else:
        studies = db.query("studies", {"project_id": project_id}, order_by="created_at DESC")

    ext_by_study: dict[str, dict[str, str]] = {}
    rob_by_study: dict[str, dict] = {}
    detailed_extractions: list[dict] = []

    if studies:
        for s in studies:
            sid = s.get("study_id", "")
            ext_by_study[sid] = {}
            if include_rob:
                rob_by_study[sid] = {}

        try:
            rows = db.execute(
                """SELECT e.*, s.study_label, s.title as study_title, s.authors as study_authors,
                          v.name as variable_name, v.section as variable_section
                   FROM extractions e
                   JOIN studies s ON e.study_id = s.study_id
                   LEFT JOIN variables v ON e.variable_id = v.variable_id
                   WHERE s.project_id = ?
                   ORDER BY v.order_index ASC, s.created_at DESC""",
                (project_id,),
            ).fetchall()
            for r in rows:
                d = dict(r)
                sid = d["study_id"]
                vid = d["variable_id"]
                if sid in ext_by_study and vid:
                    ext_by_study[sid][vid] = d.get("value") or ""
                detailed_extractions.append(d)
        except Exception:
            for s in studies:
                sid = s.get("study_id", "")
                ext_list = db.query("extractions", {"study_id": sid})
                ext_by_study[sid] = {e.get("variable_id"): e.get("value", "") for e in ext_list if e.get("variable_id")}

        if include_rob:
            rob_list = db.query("rob_assessments", {"project_id": project_id})
            for r in rob_list:
                sid = r.get("study_id")
                if sid and sid not in rob_by_study:
                    rob_by_study[sid] = r

    return proj, variables, studies, ext_by_study, rob_by_study, detailed_extractions


async def _export_csv(project_id: str, timestamp: str, include_rob: bool, db: DBBackend) -> Response:
    proj, variables, studies, ext_by_study, rob_by_study, _ = _fetch_export_data(project_id, include_rob, db)

    output = io.StringIO()
    writer = csv.writer(output)

    # Determine headers
    has_study_id = any(v.get("name", "").strip().lower() in ("study id", "study_id") for v in variables)
    headers = []
    if not has_study_id:
        headers.append("Study ID")
    for v in variables:
        headers.append(v.get("name", "").strip())
    if include_rob:
        headers += ["RoB Tool", "RoB Overall"]
    writer.writerow(headers)

    for s in studies:
        study_id = s.get("study_id", "")
        study_label = s.get("study_label") or s.get("authors", "").split(",")[0] or study_id[:8]
        row = []
        if not has_study_id:
            row.append(study_label)
        ext_map = ext_by_study.get(study_id, {})
        for v in variables:
            vname_lower = v.get("name", "").strip().lower()
            val = ext_map.get(v.get("variable_id"), "")
            if not val:
                # Pre-fill bibliographic variable if matching name
                if vname_lower in ("year", "publication year"):
                    val = str(s.get("publication_year") or "")
                elif vname_lower in ("author", "authors"):
                    val = s.get("authors", "")
                elif vname_lower in ("study id", "study_id"):
                    val = study_label
                elif vname_lower in ("doi",):
                    val = s.get("doi", "")
            row.append(_sanitize_cell_value(val))

        if include_rob:
            rob = rob_by_study.get(study_id, {})
            row.append(rob.get("tool", ""))
            row.append(rob.get("overall_judgment", ""))
        writer.writerow(row)

    csv_bytes = output.getvalue().encode("utf-8")
    filename = f"project_{project_id}_{timestamp}.csv"
    storage = get_storage("export")
    storage.save_bytes(filename, csv_bytes)

    return Response(
        content=csv_bytes,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


async def _export_xlsx(project_id: str, timestamp: str, include_rob: bool, db: DBBackend, study_id: Optional[str] = None) -> Response:
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    proj, variables, studies, ext_by_study, rob_by_study, detailed_extractions = _fetch_export_data(project_id, include_rob, db, study_id=study_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extraction Data"

    # Determine headers in exact saved sequence
    has_study_id = any(v.get("name", "").strip().lower() in ("study id", "study_id") for v in variables)
    headers = []
    if not has_study_id:
        headers.append("Study ID")
    col_vars = []
    for v in variables:
        headers.append(v.get("name", "").strip())
        col_vars.append(v)
    if include_rob:
        headers += ["RoB Tool", "RoB Overall"]
    ws.append(headers)

    # Style Header Row
    header_fill = PatternFill(start_color="3A0F34", end_color="3A0F34", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FBF9F4")
    thin_border = Border(
        left=Side(style='thin', color='D2C2CB'),
        right=Side(style='thin', color='D2C2CB'),
        top=Side(style='thin', color='D2C2CB'),
        bottom=Side(style='thin', color='D2C2CB'),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Populate Study Rows (one row per study)
    for s in studies:
        study_id = s.get("study_id", "")
        study_label = s.get("study_label") or s.get("authors", "").split(",")[0] or study_id[:8]
        row = []
        if not has_study_id:
            row.append(study_label)
        ext_map = ext_by_study.get(study_id, {})
        for v in variables:
            vname_lower = v.get("name", "").strip().lower()
            val = ext_map.get(v.get("variable_id"), "")
            if not val:
                if vname_lower in ("year", "publication year"):
                    val = str(s.get("publication_year") or "")
                elif vname_lower in ("author", "authors"):
                    val = s.get("authors", "")
                elif vname_lower in ("study id", "study_id"):
                    val = study_label
                elif vname_lower in ("doi",):
                    val = s.get("doi", "")
            row.append(_sanitize_cell_value(val))

        if include_rob:
            rob = rob_by_study.get(study_id, {})
            row.append(rob.get("tool", ""))
            row.append(rob.get("overall_judgment", ""))
        ws.append(row)

    # Freeze panes: Header row and Study ID column
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{max(2, ws.max_row)}"

    # Auto-adjust column widths and cell styling
    fill_99 = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")  # Yellow (-99)
    fill_77 = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")  # Grey (-77)
    fill_88 = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")  # Orange (-88)

    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            val_str = str(cell.value).strip() if cell.value is not None else ""
            if val_str == "-99":
                cell.fill = fill_99
            elif val_str == "-77":
                cell.fill = fill_77
            elif val_str == "-88":
                cell.fill = fill_88

    for col_idx in range(1, len(headers) + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, min(ws.max_row + 1, 20)))
        ws.column_dimensions[col_letter].width = max(16, min(max_len + 3, 50))

    # Add categorical data validations
    start_var_offset = 1 if not has_study_id else 0
    for idx, v in enumerate(col_vars):
        allowed_str = v.get("allowed_values", "")
        if allowed_str and v.get("field_type") == "categorical":
            opts = [x.strip() for x in allowed_str.replace(";", ",").split(",") if x.strip()]
            if opts:
                formula_str = f'"{",".join(opts)}"'
                dv = DataValidation(type="list", formula1=formula_str, allow_blank=True)
                ws.add_data_validation(dv)
                col_letter = openpyxl.utils.get_column_letter(start_var_offset + idx + 1)
                dv.add(f"{col_letter}2:{col_letter}{max(10, ws.max_row)}")

    # ═══════════════════════════════════════════════════════════════════
    #  Worksheet 2: Evidence (Linkage between values, quotes, and pages)
    # ═══════════════════════════════════════════════════════════════════
    ws_evidence = wb.create_sheet(title="Evidence")
    evidence_headers = [
        "Study ID", "Study Label", "Variable ID", "Variable Name",
        "Extracted Value", "Evidence Quote", "Source Document",
        "Page Number", "Confidence", "Provenance", "Verification Status"
    ]
    ws_evidence.append(evidence_headers)
    for col_idx in range(1, len(evidence_headers) + 1):
        cell = ws_evidence.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    study_label_map = {s.get("study_id"): (s.get("study_label") or s.get("title") or s.get("study_id")[:8]) for s in studies}

    for ext in detailed_extractions:
        val = ext.get("value") or ""
        quote = ext.get("source_text") or ext.get("quote") or ""
        if val or quote:
            sid = ext.get("study_id", "")
            ws_evidence.append([
                sid,
                study_label_map.get(sid, sid[:8]),
                ext.get("variable_id", ""),
                ext.get("variable_name", ""),
                _sanitize_cell_value(val),
                quote,
                ext.get("source_doc") or ext.get("study_title") or "",
                ext.get("source_page") or "",
                ext.get("confidence") if ext.get("confidence") is not None else "",
                ext.get("extracted_by") or "manual",
                ext.get("verification_status") or "pending",
            ])

    ws_evidence.freeze_panes = "A2"
    ws_evidence.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(evidence_headers))}{max(2, ws_evidence.max_row)}"

    for col_idx in range(1, len(evidence_headers) + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        max_len = max(len(str(ws_evidence.cell(row=r, column=col_idx).value or "")) for r in range(1, min(ws_evidence.max_row + 1, 25)))
        ws_evidence.column_dimensions[col_letter].width = max(14, min(max_len + 3, 45))

    # ═══════════════════════════════════════════════════════════════════
    #  Worksheet 3: Extraction Legend
    # ═══════════════════════════════════════════════════════════════════
    ws_legend = wb.create_sheet(title="Extraction_Legend")
    ws_legend.append(["Code", "Meaning", "Description"])
    ws_legend.cell(row=1, column=1).font = Font(bold=True)
    ws_legend.cell(row=1, column=2).font = Font(bold=True)
    ws_legend.cell(row=1, column=3).font = Font(bold=True)

    legend_data = [
        ("-99", "Not reported", "Searched abstract, tables, results, supplement, and derivation; value was not reported in study", fill_99),
        ("-77", "Not applicable", "Field does not apply to this study design (e.g. comparator in single-arm study)", fill_77),
        ("-88", "Unable to extract", "PDF is scanned or image-only without an extractable text layer", fill_88),
    ]

    for r_idx, (code, meaning, desc, fill) in enumerate(legend_data, start=2):
        c1 = ws_legend.cell(row=r_idx, column=1, value=code)
        c1.fill = fill
        c1.font = Font(bold=True)
        ws_legend.cell(row=r_idx, column=2, value=meaning)
        ws_legend.cell(row=r_idx, column=3, value=desc)

    out_stream = io.BytesIO()
    wb.save(out_stream)
    xlsx_bytes = out_stream.getvalue()
    filename = f"study_{study_id}_{timestamp}.xlsx" if study_id else f"project_{project_id}_{timestamp}.xlsx"
    storage = get_storage("export")
    storage.save_bytes(filename, xlsx_bytes)

    return Response(
        content=xlsx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )



@router.get("/projects/{project_id}/studies/{study_id}/export.xlsx")
async def export_single_study_xlsx(
    project_id: str,
    study_id: str,
    include_rob: bool = Query(True, description="Include RoB data in export"),
    user: dict = Depends(get_current_user_optional),
    db: DBBackend = Depends(get_db),
):
    """Export single study extraction sheet with variables and evidence linkage as formatted Excel workbook (.xlsx)."""
    check_project_access(db, project_id, user)
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    return await _export_xlsx(project_id, timestamp, include_rob, db, study_id=study_id)

# ═══════════════════════════════════════════════════════════════════════
#  Reference Manager Export (RIS, BibTeX, EndNote XML)
# ═══════════════════════════════════════════════════════════════════════

@router.get("/projects/{project_id}/export/references")
async def export_references(
    project_id: str,
    format: str = Query("ris", pattern="^(ris|bibtex|endnote)$"),
    status: str = Query("included", pattern="^(included|all|excluded)$", description="Filter by screening status: included, all, excluded"),
    user: dict = Depends(get_current_user_optional),
    db: DBBackend = Depends(get_db),
):
    """Export studies as reference manager format (RIS, BibTeX, or EndNote XML).

    By default exports only included studies. Use ?status=all for all studies.
    """
    check_project_access(db, project_id, user)

    if not db.get("projects", project_id):
        raise HTTPException(status_code=404, detail="Project not found")

    # Get studies
    if status == "all":
        studies = db.query("studies", {"project_id": project_id}, order_by="title ASC")
    elif status == "excluded":
        studies = db.query("studies", {"project_id": project_id, "screening_status": "excluded"}, order_by="title ASC")
    else:
        studies = db.query("studies", {"project_id": project_id, "screening_status": "included"}, order_by="title ASC")

    if not studies:
        raise HTTPException(status_code=404, detail="No studies found matching the criteria.")

    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")

    if format == "ris":
        content = export_ris(studies)
        ext = "ris"
        media = "application/x-research-info-systems"
    elif format == "bibtex":
        content = export_bibtex(studies)
        ext = "bib"
        media = "application/x-bibtex"
    else:  # endnote
        content = export_endnote_xml(studies)
        ext = "xml"
        media = "application/xml"

    content_bytes = content.encode("utf-8")
    filename = f"references_{project_id}_{status}_{timestamp}.{ext}"

    return Response(
        content=content_bytes,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

