"""PRISMA 2020 flow diagram export service.

Single source of truth for PRISMA flow counts (get_prisma_flow_data) plus
renderers for downloadable formats:
  - PNG  (R grid graphics png() device; cairosvg/matplotlib fallbacks)
  - DOCX (python-docx: embedded diagram + fully editable data tables)

The SVG path (r_plot_service.generate_prisma_flow) is unchanged.
"""
from __future__ import annotations

import io
import logging
from typing import Optional

from services.prisma_service import get_prisma_2020_data

logger = logging.getLogger(__name__)


def get_prisma_flow_data(project_id: str, db) -> dict:
    """Compute PRISMA 2020 flow counts for a project.

    Delegates to services.prisma_service.get_prisma_2020_data to guarantee
    live sync across databases, registers, and human/AI screening decisions.
    """
    return get_prisma_2020_data(project_id, db)


def render_prisma_png(flow_data: dict) -> Optional[bytes]:
    """Render the PRISMA diagram to PNG bytes.

    Primary: R prisma_flow.R with png() device (selected by .png extension).
    Fallback 1: convert the SVG to PNG with cairosvg.
    Fallback 2: generate diagram with matplotlib / PIL in Python.
    """
    from services.r_plot_service import call_r_script_bytes, generate_prisma_flow

    # Primary: R png() device
    try:
        png = call_r_script_bytes("prisma_flow.R", dict(flow_data), ext=".png")
        if png:
            return png
    except Exception as e:
        logger.warning(f"[PRISMA] R PNG render failed: {e}")

    # Fallback 1: SVG -> PNG via cairosvg
    try:
        import cairosvg
        svg = generate_prisma_flow(flow_data)
        if svg:
            return cairosvg.svg2png(
                bytestring=svg.encode("utf-8"),
                output_width=2200,
                output_height=1600,
                background_color="white",
            )
    except Exception as e:
        logger.warning(f"[PRISMA] cairosvg PNG fallback failed: {e}")

    # Fallback 2: Pure Python matplotlib PRISMA chart generator
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.patches import FancyBboxPatch

        fig, ax = plt.subplots(figsize=(10, 12), dpi=150)
        ax.axis("off")

        id_total = flow_data.get("identification", {}).get("total", 0)
        screened = flow_data.get("screening", {}).get("records_screened", 0)
        excluded = flow_data.get("screening", {}).get("records_excluded", 0)
        ft_assessed = flow_data.get("eligibility", {}).get("full_text_assessed", 0)
        ft_excluded = flow_data.get("eligibility", {}).get("full_text_excluded", 0)
        included = flow_data.get("included", {}).get("studies_in_review", 0)

        boxes = [
            (0.5, 0.85, f"Identification\nRecords identified: {id_total}"),
            (0.5, 0.65, f"Screening\nRecords screened: {screened}\nExcluded: {excluded}"),
            (0.5, 0.45, f"Eligibility\nFull-text assessed: {ft_assessed}\nExcluded: {ft_excluded}"),
            (0.5, 0.25, f"Included\nStudies included: {included}"),
        ]

        for x, y, text in boxes:
            box_patch = FancyBboxPatch((x - 0.3, y - 0.07), 0.6, 0.14,
                                       boxstyle="round,pad=0.03",
                                       fc="#E3F2FD", ec="#1565C0", lw=2)
            ax.add_patch(box_patch)
            ax.text(x, y, text, ha="center", va="center", fontsize=11, fontweight="bold")

        buf = io.BytesIO()
        plt.tight_layout()
        plt.savefig(buf, format="png", bbox_inches="tight")
        plt.close(fig)
        return buf.getvalue()
    except Exception as e:
        logger.warning(f"[PRISMA] matplotlib PNG fallback failed: {e}")

    return None


def _add_table(doc, title: str, rows: list[tuple[str, object]]):
    """Add a headed, editable 2-column table to the document."""
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Count"
    for i, (label, value) in enumerate(rows, start=1):
        table.rows[i].cells[0].text = str(label)
        table.rows[i].cells[1].text = str(value)


def render_prisma_docx(flow_data: dict, project_name: str = "",
                       png_bytes: Optional[bytes] = None) -> bytes:
    """Build an editable Word document for the PRISMA 2020 flow.

    Contains the diagram image (when available) followed by native Word
    tables holding every count and exclusion reason, so the numbers can be
    edited afterwards for manuscripts.
    """
    from docx import Document
    from docx.shared import Inches

    id_ = flow_data["identification"]
    scr = flow_data["screening"]
    elig = flow_data["eligibility"]
    incl = flow_data["included"]
    pending = flow_data.get("pending", 0)

    doc = Document()
    title = "PRISMA 2020 Flow Diagram"
    if project_name:
        title += f" — {project_name}"
    doc.add_heading(title, level=0)

    if png_bytes:
        doc.add_picture(io.BytesIO(png_bytes), width=Inches(6.5))
        doc.add_paragraph(
            "The diagram above is a static image. All counts below are editable "
            "tables — update the numbers there, then regenerate or adjust your "
            "diagram for publication."
        )

    _add_table(doc, "Identification", [
        ("Records identified from databases", id_["database_records"]),
        ("Records identified from registers", id_["registers"]),
        ("Records identified from other sources", id_["other_sources"]),
        ("Total records identified", id_["total"]),
    ])

    _add_table(doc, "Screening (title/abstract)", [
        ("Records screened", scr["records_screened"]),
        ("Records excluded", scr["records_excluded"]),
    ])

    ta_rows = [(r["reason"], r["count"]) for r in scr.get("excluded_with_reasons", [])]
    if ta_rows:
        _add_table(doc, "Title/abstract exclusion reasons", ta_rows)

    _add_table(doc, "Eligibility (full-text)", [
        ("Full-text articles assessed", elig["full_text_assessed"]),
        ("Full-text articles excluded", elig["full_text_excluded"]),
    ])

    ft_rows = [(r["reason"], r["count"]) for r in elig.get("excluded_reasons", [])]
    if ft_rows:
        _add_table(doc, "Full-text exclusion reasons", ft_rows)

    _add_table(doc, "Included", [
        ("Studies included in review", incl["studies_in_review"]),
        ("Studies included in meta-analysis", incl["studies_in_meta_analysis"]),
        ("Records pending screening", pending),
    ])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
