"""Coded extraction Excel workbook service.

Builds the review-facing Excel sheet where the coding is visible:
  - One sheet per variable section: StudyID rows x variable columns
  - Missing-data codes in cells with color fills:
      -99 not reported  -> yellow FFD700
      -77 not applicable -> grey  D3D3D3
      -88 unable         -> orange FFA500
      conflicts          -> light red FF6666
  - Notes column carrying the confidence tags and source citations, e.g.
      [DS] Design | [CALC] FollowUp: 2 yrs x 12 = 24 mo | [PDF-p.4: "quote"]
  - Prompts sheet: the generated extraction prompt behind every variable
  - Extraction_Legend sheet: full coding key
"""
from __future__ import annotations

import io
import logging
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

logger = logging.getLogger("radextract.extraction_excel")

# Missing-code fills (skill-file color spec)
FILL_MINUS_99 = PatternFill("solid", fgColor="FFD700")   # yellow — not reported
FILL_MINUS_77 = PatternFill("solid", fgColor="D3D3D3")   # grey — not applicable
FILL_MINUS_88 = PatternFill("solid", fgColor="FFA500")   # orange — unable to extract
FILL_CONFLICT = PatternFill("solid", fgColor="FF6666")   # light red — conflict
FILL_HEADER = PatternFill("solid", fgColor="ECE9E2")

MISSING_FILLS = {"-99": FILL_MINUS_99, "-77": FILL_MINUS_77, "-88": FILL_MINUS_88}

LEGEND_ROWS = [
    ("Code / Tag", "Meaning", "Cell color"),
    ("-99", "Not reported (exhaustive search + recheck pass completed)", "Yellow FFD700"),
    ("-77", "Not applicable to this study", "Grey D3D3D3"),
    ("-88", "Unable to extract (figure-only value or garbled PDF)", "Orange FFA500"),
    ("(conflict)", "Conflicting values found in the source", "Light red FF6666"),
    ("[DS]", "Directly stated — verbatim from table or text", ""),
    ("[INF]", "Inferred — logically derived from context", ""),
    ("[CALC]", "Calculated — arithmetic shown in the Notes column", ""),
    ("[SUP]", "From supplementary material", ""),
    ("[FIG]", "From a figure — approximated visually", ""),
    ('[PDF-p.N: "quote"]', "Source citation: page number + verbatim quote", ""),
]


def _note_for(ex: dict) -> str:
    """Build the Notes-column entry for one extraction cell."""
    parts = []
    tag = (ex.get("confidence_tag") or "").strip()
    name = ex.get("variable_name") or ""
    if tag:
        entry = f"[{tag}] {name}"
        if tag == "CALC" and ex.get("calc_note"):
            entry += f": {ex['calc_note']}"
        parts.append(entry)
    quote = (ex.get("source_text") or "").strip()
    if quote:
        page = ex.get("source_page") or 1
        parts.append(f'[PDF-p.{page}: "{quote[:120]}"]')
    return " | ".join(parts)


def build_coded_extraction_workbook(project_id: str, db) -> bytes:
    """Build the coded extraction workbook for all included studies in a project."""
    variables = db.query("variables", {"project_id": project_id}, order_by="section")
    studies = [
        s for s in db.query("studies", {"project_id": project_id}, order_by="title")
        if s.get("screening_status") == "included"
    ]

    # Group variables by section, preserving order
    sections: dict[str, list[dict]] = {}
    for v in variables:
        sections.setdefault(v.get("section") or "General", []).append(v)

    # Fetch all extractions for the included studies
    study_ids = [s["study_id"] for s in studies]
    extractions: dict[tuple[str, str], dict] = {}
    if study_ids:
        placeholders = ",".join("?" for _ in study_ids)
        rows = db.execute(
            f"""SELECT e.*, v.name as variable_name, v.section as variable_section
                FROM extractions e JOIN variables v ON e.variable_id = v.variable_id
                WHERE e.study_id IN ({placeholders})""",
            tuple(study_ids),
        ).fetchall()
        for r in rows:
            extractions[(r["study_id"], r["variable_id"])] = dict(r)

    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    # ── One sheet per section ──
    for section, sec_vars in sections.items():
        ws = wb.create_sheet(title=section[:31])
        headers = ["StudyID", "Study"] + [v["name"] for v in sec_vars] + ["Notes"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = FILL_HEADER

        for s in studies:
            row = [s["study_id"][:8], (s.get("title") or "")[:60]]
            notes = []
            for v in sec_vars:
                ex = extractions.get((s["study_id"], v["variable_id"]))
                if ex:
                    row.append(ex.get("value") or "")
                    note = _note_for(ex)
                    if note:
                        notes.append(note)
                else:
                    row.append("")
            row.append(" | ".join(notes))
            ws.append(row)

            # Color missing-code cells
            r_idx = ws.max_row
            for c_offset, v in enumerate(sec_vars, start=3):
                ex = extractions.get((s["study_id"], v["variable_id"]))
                if not ex:
                    continue
                cell = ws.cell(row=r_idx, column=c_offset)
                code = (ex.get("missing_code") or "").strip()
                if code in MISSING_FILLS:
                    cell.fill = MISSING_FILLS[code]
                if ex.get("verification_status") == "conflict":
                    cell.fill = FILL_CONFLICT

        # Column widths + wrapping
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 40
        for c_offset in range(3, 3 + len(sec_vars)):
            ws.column_dimensions[ws.cell(row=1, column=c_offset).column_letter].width = 18
        ws.column_dimensions[ws.cell(row=1, column=3 + len(sec_vars)).column_letter].width = 60
        for row_cells in ws.iter_rows(min_row=2):
            for cell in row_cells:
                cell.alignment = wrap

    # ── Prompts sheet: the instruction behind every variable ──
    ws_p = wb.create_sheet(title="Prompts")
    ws_p.append(["Section", "Variable", "Extraction prompt"])
    for cell in ws_p[1]:
        cell.font = header_font
        cell.fill = FILL_HEADER
    from services.extraction_prompt_builder import build_field_prompt
    for v in variables:
        ws_p.append([v.get("section") or "General", v["name"], build_field_prompt(v)])
    ws_p.column_dimensions["A"].width = 20
    ws_p.column_dimensions["B"].width = 28
    ws_p.column_dimensions["C"].width = 90
    for row_cells in ws_p.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = wrap

    # ── Legend sheet ──
    ws_l = wb.create_sheet(title="Extraction_Legend")
    for row in LEGEND_ROWS:
        ws_l.append(list(row))
    for cell in ws_l[1]:
        cell.font = header_font
        cell.fill = FILL_HEADER
    # Apply the actual fills to the legend color column
    legend_fills = {"-99": FILL_MINUS_99, "-77": FILL_MINUS_77, "-88": FILL_MINUS_88,
                    "(conflict)": FILL_CONFLICT}
    for r_idx in range(2, ws_l.max_row + 1):
        code = ws_l.cell(row=r_idx, column=1).value
        if code in legend_fills:
            ws_l.cell(row=r_idx, column=3).fill = legend_fills[code]
    ws_l.column_dimensions["A"].width = 22
    ws_l.column_dimensions["B"].width = 70
    ws_l.column_dimensions["C"].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
